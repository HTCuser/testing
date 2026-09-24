"""Trích xuất văn bản từ các định dạng tài liệu kỹ thuật thường gặp."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from .textutils import clean_text


class ExtractionError(RuntimeError):
    pass


def extract(path: Path) -> list[tuple[int | None, str]]:
    """Trả về danh sách (số trang, nội dung). Trang là None nếu định dạng không phân trang."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix == ".xlsx":
        return _xlsx(path)
    if suffix == ".csv":
        return _csv(path)
    if suffix in (".txt", ".md"):
        return [(None, clean_text(path.read_text(encoding="utf-8", errors="replace")))]
    raise ExtractionError(f"Chưa hỗ trợ định dạng {suffix}")


def _pdf(path: Path) -> list[tuple[int | None, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[tuple[int | None, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception:
            text = ""
        if text:
            pages.append((i, text))
    if not pages:
        raise ExtractionError(
            "Không đọc được nội dung văn bản trong PDF. "
            "Nhiều khả năng đây là bản scan ảnh, cần OCR trước khi nạp."
        )
    return pages


def _docx(path: Path) -> list[tuple[int | None, str]]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    parts: list[str] = []
    # Duyệt xen kẽ đoạn văn và bảng theo đúng thứ tự trong file. Đọc hết đoạn
    # văn rồi mới đọc bảng sẽ khiến bảng thông số mất liên hệ với tiêu đề mục
    # của nó, và khi truy hồi ra một dòng bảng thì không còn biết nó thuộc mục nào.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            parts.append(f"\n## {text}\n" if style.startswith("heading") else text)
        elif tag == "tbl":
            rows = _render_table(Table(child, doc))
            if rows:
                parts.append("\n" + "\n".join(rows) + "\n")

    body = clean_text("\n".join(parts))
    if not body:
        raise ExtractionError("Tệp DOCX không có nội dung văn bản.")
    return [(None, body)]


def _row_values(row) -> list[str]:
    return [c.text.strip().replace("\n", " ") for c in row.cells]


def _is_number(value: str) -> bool:
    return bool(value) and all(c.isdigit() or c in ",.-" for c in value)


def _header_of(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Tách tiêu đề cột, gộp lại nếu tiêu đề trải hai tầng.

    Quy trình kỹ thuật hay dùng tiêu đề hai tầng: tầng trên gộp ngang ("Các bảo
    vệ của hệ thống"), tầng dưới chia nhỏ ("Tiếng Anh", "Tiếng Việt"). Lấy mỗi
    tầng trên làm tên cột sẽ gán nhầm nhãn cho toàn bộ dữ liệu bên dưới.
    """
    top = rows[0]
    if len(rows) < 2:
        return top, []
    sub = rows[1]
    merged_top = any(top[i] and top[i] == top[i - 1] for i in range(1, len(top)))
    filled = [v for v in sub if v]
    looks_like_labels = (
        bool(filled)
        and all(len(v) <= 40 for v in filled)
        and sum(_is_number(v) for v in filled) * 2 <= len(filled)
        and any(sub[i] != top[i] for i in range(min(len(sub), len(top))))
    )
    if not (merged_top and looks_like_labels):
        return top, rows[1:]

    combined = []
    for i in range(len(top)):
        t, s = top[i], sub[i] if i < len(sub) else ""
        combined.append(f"{t} - {s}" if s and t and s != t else (s or t))
    return combined, rows[2:]


def _render_table(table) -> list[str]:
    """Kết xuất mỗi dòng bảng thành một bản ghi tự mô tả.

    Bảng trong quy trình kỹ thuật là nơi chứa thông số và cách xử lý sự cố. Nếu
    chỉ nối các ô bằng dấu "|" thì khi truy hồi ra một dòng giữa bảng, người đọc
    chỉ thấy "12,5 | m" mà không biết đó là thông số gì — dòng tiêu đề cột nằm ở
    đoạn chỉ mục khác. Ghép sẵn tên cột vào từng ô để mỗi dòng đứng một mình vẫn
    đọc được và tìm được.
    """
    rows = [_row_values(r) for r in table.rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []

    # Ô gộp bị python-docx lặp lại nguyên văn; đếm số giá trị khác nhau để biết
    # dòng đó là tiêu đề nhóm (chỉ một nội dung trải hết chiều ngang).
    if len(set(v for v in rows[0] if v)) < 2:
        header, body_rows = [], rows
    else:
        header, body_rows = _header_of(rows)

    out: list[str] = []
    group = ""
    for values in body_rows:
        distinct = [v for v in dict.fromkeys(values) if v]
        if len(distinct) == 1:
            group = distinct[0].rstrip(":")
            out.append(f"— {group}")
            continue

        fields: list[str] = []
        previous = None
        for i, value in enumerate(values):
            if not value or value == previous:
                previous = value
                continue
            previous = value
            label = header[i].strip() if i < len(header) else ""
            if label and label.lower() not in ("stt", "tt"):
                fields.append(f"{label}: {value}")
            elif not label:
                fields.append(value)
        if not fields:
            continue
        line = "; ".join(fields)
        out.append(f"[{group}] {line}" if group else line)
    return out


def _xlsx(path: Path) -> list[tuple[int | None, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[tuple[int | None, str]] = []
    for sheet in wb.worksheets:
        lines = [f"## Bảng: {sheet.title}"]
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            out.append((None, clean_text("\n".join(lines))))
    wb.close()
    if not out:
        raise ExtractionError("Bảng tính không có dữ liệu.")
    return out


def _csv(path: Path) -> list[tuple[int | None, str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    lines = [" | ".join(c.strip() for c in row) for row in reader if any(c.strip() for c in row)]
    if not lines:
        raise ExtractionError("Tệp CSV rỗng.")
    return [(None, clean_text("\n".join(lines)))]
