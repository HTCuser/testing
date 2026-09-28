"""Phiếu thao tác theo cách tổ chức của NKVH điện tử.

PTT mẫu: Nhóm → Tên phiếu → bảng bước (Mục, Địa điểm, Bước, Nội dung).
Phiếu thao tác: lập từ mẫu, cấp số tự động, đi qua các mốc Lập → Duyệt →
Tiếp nhận → Hoàn thành; vận hành viên tích từng bước ngay trên phiếu. Xuất ra
đúng mẫu phiếu Word của nhà máy để in ký.
"""
from __future__ import annotations

import io
import json
import re
import uuid
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response
from pydantic import BaseModel, Field

from .. import config, docview, phieu
from ..db import get_conn, query, query_one
from ..rag.indexer import index_record, remove_record
from ..rag.textutils import normalize, strip_accents
from ..usage import _now

router = APIRouter(prefix="/api/ptt", tags=["Phiếu thao tác"])

KINDS = {"ke_hoach": ("Kế hoạch", "KH"), "dot_xuat": ("Đột xuất", "ĐX")}
STATUSES = {
    "moi_lap": "Mới lập",
    "da_duyet": "Đã duyệt",
    "dang_thuc_hien": "Đang thực hiện",
    "hoan_thanh": "Hoàn thành",
    "huy": "Đã huỷ",
}
OPEN = ("moi_lap", "da_duyet", "dang_thuc_hien")
DEFAULT_FORMAT = "###/YYYY/{PL}/HHC"
DEFAULT_LAYOUT = Path(__file__).resolve().parent.parent / "mau" / "mau-in-phieu-thao-tac.docx"
DEFAULT_GROUPS = [
    "Vận hành bình thường - Phiếu cô lập",
    "Vận hành bình thường - Phiếu tái lập",
    "Bảo dưỡng, sửa chữa - Phiếu cô lập",
    "Bảo dưỡng, sửa chữa - Phiếu tái lập",
]


# ======================================================================= tiện ích

def _now_iso() -> str:
    return _now().strftime("%Y-%m-%dT%H:%M:%S")


def _get_setting(key: str, default: str = "") -> str:
    row = query_one("SELECT value FROM app_settings WHERE key = ?", (key,))
    return row["value"] if row else default


def _set_setting(key: str, value: str) -> None:
    conn = get_conn()
    conn.execute("INSERT INTO app_settings (key, value) VALUES (?, ?) "
                 "ON CONFLICT(key) DO UPDATE SET value = excluded.value", (key, value))
    conn.commit()


def _commit(sql: str, params: tuple = ()) -> int:
    conn = get_conn()
    cur = conn.execute(sql, params)
    conn.commit()
    return cur.lastrowid


def _loads(raw, default):
    try:
        value = json.loads(raw or "")
    except (TypeError, ValueError):
        return default
    return value if isinstance(value, type(default)) else default


def _dumps(value) -> str:
    return json.dumps(value, ensure_ascii=False)


class Step(BaseModel):
    section: str = ""
    location: str = ""
    content: str = ""


def _clean_steps(steps: list[Step]) -> list[dict]:
    """Bỏ dòng trống hoàn toàn; bước không có nội dung thì không hợp lệ."""
    out = []
    for i, s in enumerate(steps, start=1):
        row = {"section": s.section.strip(), "location": s.location.strip(), "content": s.content.strip()}
        if not any(row.values()):
            continue
        if not row["content"]:
            raise HTTPException(400, f"Dòng thứ {i} chưa có nội dung bước thao tác")
        out.append(row)
    return out


# ======================================================================= nhóm

class GroupIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)


def _ensure_default_groups() -> None:
    # Tạo sẵn bốn nhóm theo cách chia hiện tại của nhà máy, một lần duy nhất:
    # người dùng xoá hay đổi tên thì thôi, không tạo lại.
    if _get_setting("ptt_groups_seeded"):
        return
    if not query_one("SELECT id FROM ptt_groups LIMIT 1"):
        for i, name in enumerate(DEFAULT_GROUPS):
            _commit("INSERT INTO ptt_groups (name, sort_order) VALUES (?, ?)", (name, i))
    _set_setting("ptt_groups_seeded", "1")


@router.get("/nhom")
def list_groups() -> dict:
    _ensure_default_groups()
    rows = query("""SELECT g.*, (SELECT COUNT(*) FROM ptt_templates t WHERE t.group_id = g.id) AS n_templates
                      FROM ptt_groups g ORDER BY g.sort_order, g.name""")
    return {"items": [dict(r) for r in rows]}


@router.post("/nhom", status_code=201)
def create_group(payload: GroupIn) -> dict:
    top = query_one("SELECT COALESCE(MIN(sort_order), 0) - 1 AS n FROM ptt_groups")
    gid = _commit("INSERT INTO ptt_groups (name, sort_order) VALUES (?, ?)",
                  (payload.name.strip(), top["n"]))
    return dict(query_one("SELECT * FROM ptt_groups WHERE id = ?", (gid,)))


@router.put("/nhom/{group_id}")
def rename_group(group_id: int, payload: GroupIn) -> dict:
    if not query_one("SELECT id FROM ptt_groups WHERE id = ?", (group_id,)):
        raise HTTPException(404, "Không tìm thấy nhóm")
    _commit("UPDATE ptt_groups SET name = ? WHERE id = ?", (payload.name.strip(), group_id))
    for row in query("SELECT id FROM ptt_templates WHERE group_id = ?", (group_id,)):
        _index_template(row["id"])
    return dict(query_one("SELECT * FROM ptt_groups WHERE id = ?", (group_id,)))


@router.delete("/nhom/{group_id}", status_code=204)
def delete_group(group_id: int) -> Response:
    n = query_one("SELECT COUNT(*) AS n FROM ptt_templates WHERE group_id = ?", (group_id,))["n"]
    if n:
        raise HTTPException(409, f"Nhóm còn {n} phiếu mẫu. Xoá hoặc chuyển các phiếu mẫu sang nhóm khác trước.")
    _commit("DELETE FROM ptt_groups WHERE id = ?", (group_id,))
    return Response(status_code=204)


# ======================================================================= PTT mẫu

class TemplateIn(BaseModel):
    group_id: int
    name: str = Field(min_length=1, max_length=300)
    purpose: str = ""
    conditions: str = ""
    notes: str = ""
    steps: list[Step] = []


def _template(template_id: int) -> dict:
    row = query_one("""SELECT t.*, g.name AS group_name FROM ptt_templates t
                         JOIN ptt_groups g ON g.id = t.group_id WHERE t.id = ?""", (template_id,))
    if row is None:
        raise HTTPException(404, "Không tìm thấy phiếu mẫu")
    item = dict(row)
    item["steps"] = _loads(item["steps"], [])
    return item


def render_template_text(t: dict) -> str:
    """Kết xuất PTT mẫu để lập chỉ mục: trợ lý tra được trình tự thao tác chuẩn."""
    lines = [f"# Phiếu thao tác mẫu: {t['name']}", f"Nhóm: {t['group_name']}"]
    if t.get("purpose"):
        lines.append(f"Mục đích: {t['purpose']}")
    if t.get("conditions"):
        lines.append(f"\n## Điều kiện thực hiện\n{t['conditions']}")
    if t.get("notes"):
        lines.append(f"\n## Lưu ý\n{t['notes']}")
    lines.append("\n## Trình tự thao tác")
    section = location = ""
    for no, s in enumerate(t["steps"], start=1):
        section = s.get("section") or section
        location = s.get("location") or location
        where = " — ".join(x for x in (f"Mục {section}" if section else "", location) if x)
        lines.append(f"[{t['name']}] Bước {no}{f' ({where})' if where else ''}: {s['content']}")
    return "\n".join(lines)


def _index_template(template_id: int) -> None:
    t = _template(template_id)
    index_record("ptt_mau", template_id, title=f"PTT mẫu: {t['name']}", category="ptt_mau",
                 equipment_id=None, text=render_template_text(t))


@router.get("/mau")
def list_templates(nhom: int | None = None, q: str = "") -> dict:
    sql = """SELECT t.id, t.group_id, t.name, t.steps, t.updated_at, g.name AS group_name
               FROM ptt_templates t JOIN ptt_groups g ON g.id = t.group_id"""
    params: list = []
    if nhom:
        sql += " WHERE t.group_id = ?"
        params.append(nhom)
    sql += " ORDER BY t.name"
    needle = normalize(q.strip())
    items = []
    for row in query(sql, params):
        item = dict(row)
        item["n_steps"] = len(_loads(item.pop("steps"), []))
        if needle and needle not in normalize(f"{item['name']} {item['group_name']}"):
            continue
        items.append(item)
    return {"items": items, "total": len(items)}


@router.get("/mau/{template_id}")
def get_template(template_id: int) -> dict:
    return _template(template_id)


@router.post("/mau", status_code=201)
def create_template(payload: TemplateIn) -> dict:
    if not query_one("SELECT id FROM ptt_groups WHERE id = ?", (payload.group_id,)):
        raise HTTPException(400, "Nhóm không tồn tại")
    tid = _commit(
        """INSERT INTO ptt_templates (group_id, name, purpose, conditions, notes, steps)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (payload.group_id, payload.name.strip(), payload.purpose.strip(), payload.conditions.strip(),
         payload.notes.strip(), _dumps(_clean_steps(payload.steps))),
    )
    _index_template(tid)
    return _template(tid)


@router.put("/mau/{template_id}")
def update_template(template_id: int, payload: TemplateIn) -> dict:
    _template(template_id)
    _commit(
        """UPDATE ptt_templates SET group_id = ?, name = ?, purpose = ?, conditions = ?, notes = ?,
                  steps = ?, updated_at = datetime('now') WHERE id = ?""",
        (payload.group_id, payload.name.strip(), payload.purpose.strip(), payload.conditions.strip(),
         payload.notes.strip(), _dumps(_clean_steps(payload.steps)), template_id),
    )
    _index_template(template_id)
    return _template(template_id)


@router.delete("/mau/{template_id}", status_code=204)
def delete_template(template_id: int) -> Response:
    _template(template_id)
    # Phiếu đã lập giữ nguyên bản sao các bước của nó, không phụ thuộc mẫu.
    _commit("DELETE FROM ptt_templates WHERE id = ?", (template_id,))
    remove_record("ptt_mau", template_id)
    return Response(status_code=204)


@router.post("/mau/{template_id}/sao-chep", status_code=201)
def copy_template(template_id: int) -> dict:
    t = _template(template_id)
    tid = _commit(
        """INSERT INTO ptt_templates (group_id, name, purpose, conditions, notes, steps)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (t["group_id"], f"{t['name']} (bản sao)", t["purpose"], t["conditions"], t["notes"],
         _dumps(t["steps"])),
    )
    _index_template(tid)
    return _template(tid)


_COLS = {"muc": "section", "dia diem": "location", "noi dung": "content",
         "noi dung thao tac": "content", "buoc": "no", "stt": "no"}


@router.post("/nhap-excel")
async def import_excel(file: UploadFile = File(...)) -> dict:
    """Đọc bảng bước từ Excel (cột Mục, Địa điểm, Bước, Nội dung), chưa ghi gì.

    Tìm hàng tiêu đề trong 10 hàng đầu; không có tiêu đề thì hiểu 4 cột đầu
    theo thứ tự Mục, Địa điểm, Bước, Nội dung như bảng trên NKVH.
    """
    if Path(file.filename or "").suffix.lower() not in (".xlsx", ".xlsm"):
        raise HTTPException(400, "Cần file Excel .xlsx (file .xls đời cũ: mở bằng Excel rồi lưu thành .xlsx)")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Không đọc được file Excel: {exc}")
    rows = [["" if v is None else str(v).strip() for v in r]
            for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    rows = [r for r in rows if any(r)]

    mapping, start = {}, 0
    for i, row in enumerate(rows[:10]):
        found = {j: _COLS[normalize(v).strip().rstrip(":")] for j, v in enumerate(row)
                 if normalize(v).strip().rstrip(":") in _COLS}
        if "content" in found.values():
            mapping, start = found, i + 1
            break
    if not mapping:
        mapping = {0: "section", 1: "location", 2: "no", 3: "content"}

    steps = []
    for row in rows[start:]:
        step = {"section": "", "location": "", "content": ""}
        for j, key in mapping.items():
            if key != "no" and j < len(row):
                step[key] = row[j]
        if step["content"]:
            steps.append(step)
    if not steps:
        raise HTTPException(400, "Không tìm thấy bước nào. Cần một cột tiêu đề là \"Nội dung\".")
    return {"steps": steps, "total": len(steps)}


# ======================================================================= số phiếu

def _format() -> str:
    return _get_setting("ptt_number_format", DEFAULT_FORMAT)


def _shared() -> bool:
    return _get_setting("ptt_shared_sequence", "1") == "1"


def _book(kind: str) -> str:
    # Một dãy số là một "sổ". Tiền tố riêng để không lẫn với sổ của module cũ.
    fmt = _format()
    return f"ptt:{fmt}" if _shared() else f"ptt:{fmt}|{kind}"


def _code(number: int, year: int, kind: str) -> str:
    return phieu.format_number(_format(), number, year).replace("{PL}", KINDS[kind][1])


def _next_number(book: str, year: int) -> int:
    row = query_one("SELECT next_number FROM ticket_counters WHERE book = ? AND year = ?", (book, year))
    if row:
        return row["next_number"]
    used = query_one("SELECT MAX(number) AS n FROM ptt_tickets WHERE book = ? AND year = ?", (book, year))
    return (used["n"] or 0) + 1


class ConfigIn(BaseModel):
    number_format: str | None = None
    shared_sequence: bool | None = None
    next_number: int | None = None


def _config() -> dict:
    year = _now().year
    books = {k: _next_number(_book(k), year) for k in KINDS}
    layout = _get_setting("ptt_layout_name")
    return {
        "number_format": _format(),
        "shared_sequence": _shared(),
        "year": year,
        "next": {k: {"number": n, "code": _code(n, year, k)} for k, n in books.items()},
        "layout": layout or "Mẫu phiếu thao tác Hủa Na (mặc định)",
        "layout_custom": bool(layout),
        "issuing_unit": config.ORG_UNIT,
        "kinds": [{"value": k, "label": v[0], "code": v[1]} for k, v in KINDS.items()],
        "statuses": [{"value": k, "label": v} for k, v in STATUSES.items()],
        "placeholders": LAYOUT_FIELDS,
    }


@router.get("/cau-hinh")
def get_config() -> dict:
    return _config()


@router.put("/cau-hinh")
def update_config(payload: ConfigIn) -> dict:
    if payload.number_format is not None:
        fmt = payload.number_format.strip()
        if "#" not in fmt:
            raise HTTPException(400, 'Định dạng số phiếu phải có dấu "#" đánh dấu chỗ đặt số thứ tự, '
                                     "ví dụ ###/YYYY/{PL}/HHC")
        _set_setting("ptt_number_format", fmt[:60])
    if payload.shared_sequence is not None:
        _set_setting("ptt_shared_sequence", "1" if payload.shared_sequence else "0")
    if payload.next_number is not None:
        year = _now().year
        for kind in (KINDS if not _shared() else ["ke_hoach"]):
            book = _book(kind)
            used = query_one("SELECT MAX(number) AS n FROM ptt_tickets WHERE book = ? AND year = ?",
                             (book, year))["n"] or 0
            if payload.next_number <= used:
                raise HTTPException(400, f"Năm {year} đã cấp tới số {used}, số tiếp theo phải lớn hơn {used}.")
            _commit("""INSERT INTO ticket_counters (book, year, next_number) VALUES (?, ?, ?)
                       ON CONFLICT(book, year) DO UPDATE SET next_number = excluded.next_number""",
                    (book, year, payload.next_number))
    return _config()


def _layout_path() -> Path:
    name = _get_setting("ptt_layout_file")
    if name:
        path = config.TICKET_TEMPLATE_DIR / name
        if path.exists():
            return path
    return DEFAULT_LAYOUT


@router.post("/cau-hinh/mau-in")
async def upload_layout(file: UploadFile = File(...)) -> dict:
    if Path(file.filename or "").suffix.lower() != ".docx":
        raise HTTPException(400, "Mẫu in phải là file Word .docx")
    stored = f"mau-in-{uuid.uuid4().hex}.docx"
    path = config.TICKET_TEMPLATE_DIR / stored
    path.write_bytes(await file.read())
    try:
        fields = phieu.find_fields(path)
    except Exception as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(400, f"Không đọc được file Word: {exc}")
    if not fields:
        path.unlink(missing_ok=True)
        raise HTTPException(400, "Mẫu in không có ô {{...}} nào. Tải mẫu mặc định về xem cách đặt ô.")
    _set_setting("ptt_layout_file", stored)
    _set_setting("ptt_layout_name", file.filename or stored)
    return _config()


@router.delete("/cau-hinh/mau-in")
def reset_layout() -> dict:
    _set_setting("ptt_layout_file", "")
    _set_setting("ptt_layout_name", "")
    return _config()


@router.get("/cau-hinh/mau-in")
def download_layout():
    return FileResponse(_layout_path(), filename="Mau in phieu thao tac.docx")


# ======================================================================= phiếu

class Person(BaseModel):
    name: str = ""
    title: str = ""


class People(BaseModel):
    viet: Person = Person()
    duyet: Person = Person()
    giam_sat: list[Person] = []
    thao_tac: list[Person] = []


class TicketIn(BaseModel):
    template_id: int | None = None
    kind: str = "ke_hoach"
    name: str = Field(min_length=1, max_length=300)
    purpose: str = ""
    conditions: str = ""
    notes: str = ""
    requesting_unit: str = ""
    planned_start: str = ""
    planned_end: str = ""
    people: People = People()
    steps: list[Step] = []


class HandoverRow(BaseModel):
    time: str = ""
    unit: str = ""
    name: str = ""
    content: str = ""


def _clean_handover(rows: list[HandoverRow]) -> list[dict]:
    return [r for r in (x.model_dump() for x in rows) if any(str(v).strip() for v in r.values())]


class TicketUpdate(BaseModel):
    kind: str | None = None
    name: str | None = None
    purpose: str | None = None
    conditions: str | None = None
    notes: str | None = None
    requesting_unit: str | None = None
    planned_start: str | None = None
    planned_end: str | None = None
    people: People | None = None
    steps: list[Step] | None = None
    abnormal: str | None = None
    handover_before: list[HandoverRow] | None = None
    handover_after: list[HandoverRow] | None = None


def _ticket(ticket_id: int, with_steps: bool = True) -> dict:
    row = query_one("SELECT * FROM ptt_tickets WHERE id = ?", (ticket_id,))
    if row is None:
        raise HTTPException(404, "Không tìm thấy phiếu thao tác")
    t = dict(row)
    t["people"] = _loads(t["people"], {})
    t["handover_before"] = _loads(t.get("handover_before"), [])
    t["handover_after"] = _loads(t.get("handover_after"), [])
    t["status_label"] = STATUSES.get(t["status"], t["status"])
    t["kind_label"] = KINDS.get(t["kind"], (t["kind"],))[0]
    if with_steps:
        t["steps"] = [dict(r) for r in query(
            "SELECT * FROM ptt_ticket_steps WHERE ticket_id = ? ORDER BY ord", (ticket_id,))]
        t["attachments"] = [dict(r) for r in query(
            "SELECT id, filename, size_bytes, created_at FROM ptt_attachments WHERE ticket_id = ? ORDER BY id",
            (ticket_id,))]
    return t


def _replace_steps(conn, ticket_id: int, steps: list[dict]) -> None:
    conn.execute("DELETE FROM ptt_ticket_steps WHERE ticket_id = ?", (ticket_id,))
    conn.executemany(
        "INSERT INTO ptt_ticket_steps (ticket_id, ord, section, location, content) VALUES (?, ?, ?, ?, ?)",
        [(ticket_id, i, s["section"], s["location"], s["content"]) for i, s in enumerate(steps, start=1)],
    )


def _check_kind(kind: str) -> str:
    if kind not in KINDS:
        raise HTTPException(400, "Phân loại phiếu phải là Kế hoạch hoặc Đột xuất")
    return kind


@router.get("/phieu")
def list_tickets(q: str = "", status: str = "", kind: str = "", year: int | None = None) -> dict:
    sql = """SELECT t.*, (SELECT COUNT(*) FROM ptt_ticket_steps s WHERE s.ticket_id = t.id) AS n_steps,
                    (SELECT COUNT(*) FROM ptt_ticket_steps s WHERE s.ticket_id = t.id AND s.done) AS n_done
               FROM ptt_tickets t WHERE 1 = 1"""
    params: list = []
    for col, val in (("status", status), ("kind", kind)):
        if val:
            sql += f" AND t.{col} = ?"
            params.append(val)
    if year:
        sql += " AND t.year = ?"
        params.append(year)
    sql += " ORDER BY t.year DESC, t.number DESC"
    needle = normalize(q.strip())
    items = []
    for row in query(sql, params):
        t = dict(row)
        t["people"] = _loads(t["people"], {})
        t["status_label"] = STATUSES.get(t["status"], t["status"])
        t["kind_label"] = KINDS.get(t["kind"], (t["kind"],))[0]
        if needle:
            names = json.dumps(t["people"], ensure_ascii=False)
            if needle not in normalize(f"{t['code']} {t['name']} {t['purpose']} {names}"):
                continue
        items.append(t)
    years = [r["year"] for r in query("SELECT DISTINCT year FROM ptt_tickets ORDER BY year DESC")]
    return {"items": items, "total": len(items), "years": years}


@router.get("/phieu/so-tiep-theo")
def next_code(kind: str = "ke_hoach") -> dict:
    _check_kind(kind)
    year = _now().year
    n = _next_number(_book(kind), year)
    return {"number": n, "code": _code(n, year, kind)}


@router.get("/phieu/tong-hop")
def daily_summary(ngay: str = "") -> dict:
    """Kiểm soát phiếu theo ngày: phiếu trong ngày theo thời gian bắt đầu dự
    kiến, và phiếu tồn — quá ngày mà chưa hoàn thành, chưa huỷ."""
    from datetime import date, timedelta

    today = _now().date()
    try:
        day = date.fromisoformat(ngay) if ngay else today
    except ValueError:
        raise HTTPException(400, "Ngày không hợp lệ, cần dạng YYYY-MM-DD")
    iso = day.isoformat()
    day_expr = "substr(CASE WHEN planned_start <> '' THEN planned_start ELSE created_at END, 1, 10)"

    def rows(where: str, params: tuple) -> list[dict]:
        out = []
        for r in query(f"""SELECT t.*, (SELECT COUNT(*) FROM ptt_ticket_steps s WHERE s.ticket_id = t.id) AS n_steps,
                                  (SELECT COUNT(*) FROM ptt_ticket_steps s WHERE s.ticket_id = t.id AND s.done) AS n_done
                             FROM ptt_tickets t WHERE {where} ORDER BY t.planned_start, t.number""", params):
            t = dict(r)
            t["people"] = _loads(t["people"], {})
            t["status_label"] = STATUSES.get(t["status"], t["status"])
            t["kind_label"] = KINDS.get(t["kind"], (t["kind"],))[0]
            out.append(t)
        return out

    day_items = rows(f"{day_expr} = ?", (iso,))
    backlog = rows(f"{day_expr} < ? AND status IN ('moi_lap', 'da_duyet', 'dang_thuc_hien')", (iso,))
    counts = {k: 0 for k in STATUSES}
    for t in day_items:
        counts[t["status"]] += 1
    by_operator: dict[str, int] = {}
    for t in day_items:
        if t["status"] == "huy":
            continue
        for p in t["people"].get("thao_tac", []):
            if p.get("name"):
                by_operator[p["name"]] = by_operator.get(p["name"], 0) + 1
    week = []
    for offset in range(6, -1, -1):
        d = (day - timedelta(days=offset)).isoformat()
        r = query_one(f"""SELECT COUNT(*) AS lap, SUM(status = 'hoan_thanh') AS xong, SUM(status = 'huy') AS huy
                            FROM ptt_tickets WHERE {day_expr} = ?""", (d,))
        week.append({"date": d, "lap": r["lap"] or 0, "xong": r["xong"] or 0, "huy": r["huy"] or 0})
    return {"date": iso, "is_today": day == today, "counts": {**counts, "tong": len(day_items)},
            "items": day_items, "backlog": backlog,
            "by_operator": sorted(by_operator.items(), key=lambda kv: -kv[1]), "week": week}


@router.post("/phieu", status_code=201)
def create_ticket(payload: TicketIn) -> dict:
    kind = _check_kind(payload.kind)
    steps = _clean_steps(payload.steps)
    if not steps:
        raise HTTPException(400, "Phiếu chưa có bước thao tác nào")
    year = _now().year
    book = _book(kind)
    conn = get_conn()
    if conn.in_transaction:
        conn.commit()
    # Khoá ghi ngay từ đầu: hai người lập phiếu cùng lúc vẫn nhận hai số khác nhau.
    conn.execute("BEGIN IMMEDIATE")
    try:
        number = _next_number(book, year)
        cur = conn.execute(
            """INSERT INTO ptt_tickets (template_id, book, year, number, code, kind, name, purpose,
                   conditions, notes, requesting_unit, issuing_unit, planned_start, planned_end, people,
                   created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (payload.template_id, book, year, number, _code(number, year, kind), kind,
             payload.name.strip(), payload.purpose.strip(), payload.conditions.strip(),
             payload.notes.strip(), payload.requesting_unit.strip(), config.ORG_UNIT,
             payload.planned_start, payload.planned_end, _dumps(payload.people.model_dump()),
             # Giờ nhà máy, không phải giờ UTC mặc định của SQLite: ngày lập phiếu
             # in trên phiếu và đếm trên dashboard phải đúng ngày trực ca.
             _now_iso()),
        )
        ticket_id = cur.lastrowid
        _replace_steps(conn, ticket_id, steps)
        conn.execute("""INSERT INTO ticket_counters (book, year, next_number) VALUES (?, ?, ?)
                        ON CONFLICT(book, year) DO UPDATE SET next_number = excluded.next_number""",
                     (book, year, number + 1))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return _ticket(ticket_id)


@router.get("/phieu/{ticket_id}")
def get_ticket(ticket_id: int) -> dict:
    return _ticket(ticket_id)


@router.put("/phieu/{ticket_id}")
def update_ticket(ticket_id: int, payload: TicketUpdate) -> dict:
    t = _ticket(ticket_id, with_steps=False)
    if t["status"] in ("hoan_thanh", "huy"):
        raise HTTPException(409, f"Phiếu {t['status_label'].lower()}, không sửa được nữa")
    data = payload.model_dump(exclude_none=True)
    header = {"kind", "name", "purpose", "conditions", "notes", "requesting_unit",
              "planned_start", "planned_end", "people", "steps"}
    if t["status"] != "moi_lap" and header & data.keys():
        # Phiếu đã duyệt là phiếu người duyệt đã ký: nội dung phải giữ nguyên.
        raise HTTPException(409, "Phiếu đã duyệt, không sửa được nội dung. Chỉ ghi được sự kiện bất thường.")
    conn = get_conn()
    sets, params = [], []
    for key in ("name", "purpose", "conditions", "notes", "requesting_unit", "planned_start",
                "planned_end", "abnormal"):
        if key in data:
            sets.append(f"{key} = ?")
            params.append(str(data[key]).strip())
    if "kind" in data:
        # Phân loại đi kèm số phiếu (KH/ĐX): đổi phân loại thì đổi luôn chữ trong số.
        kind = _check_kind(data["kind"])
        sets += ["kind = ?", "code = ?"]
        params += [kind, _code(t["number"], t["year"], kind)]
    if "people" in data:
        sets.append("people = ?")
        params.append(_dumps(data["people"]))
    # Giao nhận, nghiệm thu diễn ra lúc thực hiện nên ghi được sau khi duyệt.
    for key in ("handover_before", "handover_after"):
        rows = getattr(payload, key)
        if rows is not None:
            sets.append(f"{key} = ?")
            params.append(_dumps(_clean_handover(rows)))
    if sets:
        conn.execute(f"UPDATE ptt_tickets SET {', '.join(sets)}, updated_at = datetime('now') WHERE id = ?",
                     (*params, ticket_id))
    if payload.steps is not None:
        steps = _clean_steps(payload.steps)
        if not steps:
            raise HTTPException(400, "Phiếu phải có ít nhất một bước thao tác")
        _replace_steps(conn, ticket_id, steps)
    conn.commit()
    return _ticket(ticket_id)


def _transition(ticket_id: int, allowed: tuple, status: str, stamp: str, extra: dict | None = None) -> dict:
    t = _ticket(ticket_id, with_steps=False)
    if t["status"] not in allowed:
        raise HTTPException(409, f"Phiếu đang ở trạng thái \"{t['status_label']}\", không thực hiện được thao tác này")
    fields = {"status": status, stamp: _now_iso(), **(extra or {})}
    _commit(f"UPDATE ptt_tickets SET {', '.join(f'{k} = ?' for k in fields)}, updated_at = datetime('now') "
            "WHERE id = ?", (*fields.values(), ticket_id))
    return _ticket(ticket_id)


@router.post("/phieu/{ticket_id}/duyet")
def approve(ticket_id: int) -> dict:
    if not query_one("SELECT id FROM ptt_ticket_steps WHERE ticket_id = ?", (ticket_id,)):
        raise HTTPException(400, "Phiếu chưa có bước thao tác nào")
    return _transition(ticket_id, ("moi_lap",), "da_duyet", "approved_at")


@router.post("/phieu/{ticket_id}/tiep-nhan")
def receive(ticket_id: int) -> dict:
    return _transition(ticket_id, ("da_duyet",), "dang_thuc_hien", "received_at")


@router.post("/phieu/{ticket_id}/hoan-thanh")
def complete(ticket_id: int) -> dict:
    return _transition(ticket_id, ("dang_thuc_hien",), "hoan_thanh", "completed_at")


class CancelIn(BaseModel):
    reason: str = ""


@router.post("/phieu/{ticket_id}/huy")
def cancel(ticket_id: int, payload: CancelIn) -> dict:
    return _transition(ticket_id, OPEN, "huy", "cancelled_at", {"cancel_reason": payload.reason.strip()})


class StepToggle(BaseModel):
    done: bool
    # Người ra lệnh / nhận lệnh bước này. Bỏ trống thì lấy người giám sát và
    # người thao tác đầu tiên của phiếu, như NKVH.
    commander: str | None = None
    receiver: str | None = None


@router.put("/phieu/{ticket_id}/buoc/{step_id}")
def toggle_step(ticket_id: int, step_id: int, payload: StepToggle) -> dict:
    """Vận hành viên tích bước đã thực hiện, ghi lại giờ tích."""
    t = _ticket(ticket_id, with_steps=False)
    if t["status"] not in ("da_duyet", "dang_thuc_hien"):
        raise HTTPException(409, "Chỉ tích được bước khi phiếu đã duyệt và chưa hoàn thành")
    step = query_one("SELECT id FROM ptt_ticket_steps WHERE id = ? AND ticket_id = ?", (step_id, ticket_id))
    if step is None:
        raise HTTPException(404, "Không tìm thấy bước")
    conn = get_conn()
    if payload.done:
        people = t["people"]
        commander = payload.commander if payload.commander is not None else \
            ((people.get("giam_sat") or [{}])[0].get("name", ""))
        receiver = payload.receiver if payload.receiver is not None else \
            ((people.get("thao_tac") or [{}])[0].get("name", ""))
        conn.execute("UPDATE ptt_ticket_steps SET done = 1, done_at = ?, commander = ?, receiver = ? WHERE id = ?",
                     (_now_iso(), commander.strip(), receiver.strip(), step_id))
    else:
        conn.execute("UPDATE ptt_ticket_steps SET done = 0, done_at = '', commander = '', receiver = '' "
                     "WHERE id = ?", (step_id,))
    if payload.done and t["status"] == "da_duyet":
        # Bắt đầu tích bước tức là đã tiếp nhận phiếu và bắt tay thao tác.
        conn.execute("UPDATE ptt_tickets SET status = 'dang_thuc_hien', received_at = ? WHERE id = ?",
                     (_now_iso(), ticket_id))
    conn.execute("UPDATE ptt_tickets SET updated_at = datetime('now') WHERE id = ?", (ticket_id,))
    conn.commit()
    return _ticket(ticket_id)


# ------------------------------------------------------------------ đính kèm

@router.post("/phieu/{ticket_id}/dinh-kem", status_code=201)
async def attach(ticket_id: int, file: UploadFile = File(...)) -> dict:
    _ticket(ticket_id, with_steps=False)
    data = await file.read()
    if len(data) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(413, f"Tệp vượt quá {config.MAX_UPLOAD_MB} MB")
    suffix = Path(file.filename or "").suffix.lower()[:10]
    stored = f"{uuid.uuid4().hex}{suffix}"
    (config.PTT_FILE_DIR / stored).write_bytes(data)
    _commit("INSERT INTO ptt_attachments (ticket_id, filename, stored_name, size_bytes) VALUES (?, ?, ?, ?)",
            (ticket_id, file.filename or stored, stored, len(data)))
    return _ticket(ticket_id)


@router.get("/dinh-kem/{attachment_id}")
def download_attachment(attachment_id: int):
    row = query_one("SELECT * FROM ptt_attachments WHERE id = ?", (attachment_id,))
    if row is None or not (config.PTT_FILE_DIR / row["stored_name"]).exists():
        raise HTTPException(404, "Không tìm thấy tệp đính kèm")
    return FileResponse(config.PTT_FILE_DIR / row["stored_name"], filename=row["filename"])


@router.delete("/dinh-kem/{attachment_id}", status_code=204)
def delete_attachment(attachment_id: int) -> Response:
    row = query_one("SELECT * FROM ptt_attachments WHERE id = ?", (attachment_id,))
    if row:
        _commit("DELETE FROM ptt_attachments WHERE id = ?", (attachment_id,))
        (config.PTT_FILE_DIR / row["stored_name"]).unlink(missing_ok=True)
    return Response(status_code=204)


# ------------------------------------------------------------------ xuất Word

# Các ô phần mềm điền vào mẫu in. Mẫu in tự làm chỉ cần đặt đúng tên các ô này.
LAYOUT_FIELDS = [
    "Số phiếu", "Tên phiếu", "Phân loại", "Mục đích", "Điều kiện", "Lưu ý", "Đơn vị đề nghị",
    "Đơn vị cấp phiếu", "Người viết phiếu", "Chức vụ người viết", "Người duyệt phiếu",
    "Chức vụ người duyệt", "Người giám sát", "Chức vụ người giám sát", "Người thao tác",
    "Chức vụ người thao tác", "Giờ bắt đầu", "Ngày bắt đầu", "Tháng bắt đầu", "Năm bắt đầu",
    "Giờ kết thúc", "Ngày kết thúc", "Tháng kết thúc", "Năm kết thúc", "Ngày", "Tháng", "Năm",
    "Sự kiện bất thường",
]


def _numbered(text: str) -> str:
    """Điều kiện, lưu ý: mỗi dòng một mục, đánh số như phiếu giấy nếu chưa có số."""
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    if not lines or all(re.match(r"^\d+\s*[.)]", ln) for ln in lines):
        return "\n".join(lines)
    return "\n".join(f"{i}. {ln}" for i, ln in enumerate(lines, start=1))


def _time_parts(value: str) -> dict:
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:T(\d{2}):(\d{2}))?", value or "")
    if not m:
        return {"gio": "", "ngay": "", "thang": "", "nam": ""}
    y, mo, d, hh, mm = m.groups()
    return {"gio": f"{hh}h{mm}" if hh else "", "ngay": str(int(d)), "thang": str(int(mo)), "nam": y}


def _layout_values(t: dict) -> dict[str, str]:
    people = t["people"]
    def person(role, i=None):
        p = people.get(role) or ({} if i is None else [])
        if i is not None:
            p = p[i] if len(p) > i else {}
        return p.get("name", ""), p.get("title", "")
    values = {
        "Số phiếu": t["code"], "Tên phiếu": t["name"], "Phân loại": t["kind_label"],
        "Mục đích": t["purpose"], "Điều kiện": _numbered(t["conditions"]), "Lưu ý": _numbered(t["notes"]),
        "Đơn vị đề nghị": t["requesting_unit"], "Đơn vị cấp phiếu": t["issuing_unit"],
        "Sự kiện bất thường": t["abnormal"],
    }
    values["Người viết phiếu"], values["Chức vụ người viết"] = person("viet")
    values["Người duyệt phiếu"], values["Chức vụ người duyệt"] = person("duyet")
    for role, label, short in (("giam_sat", "Người giám sát", "người giám sát"),
                               ("thao_tac", "Người thao tác", "người thao tác")):
        for i in range(3):
            suffix = "" if i == 0 else f" {i + 1}"
            values[f"{label}{suffix}"], values[f"Chức vụ {short}{suffix}"] = person(role, i)
    for key, label in (("planned_start", "bắt đầu"), ("planned_end", "kết thúc")):
        parts = _time_parts(t[key])
        values[f"Giờ {label}"] = parts["gio"]
        values[f"Ngày {label}"], values[f"Tháng {label}"], values[f"Năm {label}"] = \
            parts["ngay"], parts["thang"], parts["nam"]
    created = _time_parts(t["created_at"].replace(" ", "T"))
    values["Ngày"], values["Tháng"], values["Năm"] = created["ngay"], created["thang"], created["nam"]
    return values


def _render_word(ticket_id: int) -> tuple[dict, bytes]:
    t = _ticket(ticket_id)
    data = phieu.render_ticket(_layout_path(), _layout_values(t), t["steps"],
                               handover=(t["handover_before"], t["handover_after"]))
    return t, data


def _file_name(t: dict) -> str:
    name = re.sub(r'[\\/:*?"<>|]+', "-", f"PTT {t['code']} - {t['name']}")[:150]
    # Tên không dấu: có trình duyệt, máy in mạng đọc sai tên file có dấu.
    return strip_accents(name).strip() + ".docx"


@router.get("/phieu/{ticket_id}/word")
def download_word(ticket_id: int):
    t, data = _render_word(ticket_id)
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{_file_name(t)}"'},
    )


@router.get("/phieu/{ticket_id}/xem", response_class=HTMLResponse)
def view_ticket(ticket_id: int):
    t, data = _render_word(ticket_id)
    path = config.TICKET_DIR / f"ptt-{ticket_id}.docx"
    path.write_bytes(data)
    body = docview.render_body(path)
    if t["status"] == "huy":
        body = ('<p style="color:#b91c1c;font-weight:700;border:2px solid #b91c1c;padding:8px 12px;'
                'display:inline-block">PHIẾU ĐÃ HUỶ</p>' + body)
    return HTMLResponse(docview.page(f"Phiếu thao tác số {t['code']}", f"{t['name']} · {t['status_label']}",
                                     body, f"/api/ptt/phieu/{ticket_id}/word"))


# ------------------------------------------------------------------ gợi ý nhập liệu

@router.get("/goi-y")
def suggestions() -> dict:
    """Tên người, chức vụ, đơn vị đã từng nhập — chọn lại cho nhanh."""
    names, titles, units, locations = {}, {}, {}, {}
    for row in query("SELECT people, requesting_unit FROM ptt_tickets ORDER BY id DESC LIMIT 300"):
        people = _loads(row["people"], {})
        persons = [people.get("viet") or {}, people.get("duyet") or {},
                   *(people.get("giam_sat") or []), *(people.get("thao_tac") or [])]
        for p in persons:
            if p.get("name"):
                names.setdefault(p["name"], None)
            if p.get("title"):
                titles.setdefault(p["title"], None)
        if row["requesting_unit"]:
            units.setdefault(row["requesting_unit"], None)
    for row in query("SELECT steps FROM ptt_templates ORDER BY updated_at DESC LIMIT 200"):
        for s in _loads(row["steps"], []):
            if s.get("location"):
                locations.setdefault(s["location"], None)
    return {"names": list(names)[:40], "titles": list(titles)[:20] or
            ["Trưởng ca", "Trưởng kíp", "Vận hành viên", "Trực chính", "Trực phụ"],
            "units": list(units)[:20], "locations": list(locations)[:40]}
