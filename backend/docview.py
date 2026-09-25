"""Dựng bản xem tài liệu ngay trên trình duyệt.

Vận hành viên đang tra cứu thì cần mở nguyên văn quy trình để đối chiếu, chứ
không phải tải tệp về rồi mở bằng Word — trên máy trực ca nhiều khi không có
Word, và tải về mỗi lần cũng làm rơi mất ngữ cảnh đang tra. Trình duyệt đọc
được sẵn PDF và văn bản thuần; DOCX và bảng tính thì phải chuyển sang HTML.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from . import config

# Định dạng trình duyệt tự mở được, không cần chuyển đổi.
NATIVE_SUFFIXES = {".pdf", ".txt", ".md"}

# Chuyển một quy trình 2 MB mất khoảng sáu giây — mở lần nào cũng chờ chừng đó
# thì không ai dùng. Tệp gốc không đổi nên bản dựng cũng không đổi: giữ lại trên
# đĩa, khoá theo thời điểm sửa và kích thước tệp để tự hỏng khi tài liệu được
# nạp lại.
CACHE_DIR = config.INDEX_DIR / "xem"


def needs_conversion(path: Path) -> bool:
    return path.suffix.lower() not in NATIVE_SUFFIXES


def warm(path: Path) -> None:
    """Dựng sẵn bản xem ngay khi nạp tài liệu.

    Dựng lười thì lần mở đầu tiên người dùng phải nhìn tab trắng vài giây. Nạp
    tài liệu vốn đã là thao tác chờ, gộp luôn vào đó thì mọi lần mở sau đều tức thì.
    """
    if not needs_conversion(path):
        return
    try:
        render_body_cached(path)
    except Exception:
        # Bản xem hỏng không được làm hỏng việc lập chỉ mục; lúc mở sẽ báo lỗi.
        pass


def render_body_cached(path: Path) -> str:
    stat = path.stat()
    current = CACHE_DIR / f"{path.stem}-{int(stat.st_mtime)}-{stat.st_size}.html"
    if current.exists():
        return current.read_text(encoding="utf-8")

    body = render_body(path)
    try:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        for stale in CACHE_DIR.glob(f"{path.stem}-*.html"):
            if stale != current:
                stale.unlink(missing_ok=True)
        current.write_text(body, encoding="utf-8")
    except OSError:
        # Không ghi được bộ nhớ đệm thì vẫn phải trả được bản xem.
        pass
    return body


def render_body(path: Path) -> str:
    """Phần thân HTML của tài liệu. Ném ValueError nếu không dựng được."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_html(path)
    if suffix == ".xlsx":
        return _xlsx_html(path)
    if suffix == ".csv":
        return _csv_html(path)
    raise ValueError(f"Chưa hỗ trợ xem trực tiếp định dạng {suffix}")


def _docx_html(path: Path) -> str:
    try:
        import mammoth
    except ImportError:
        # Thiếu mammoth thì vẫn phải mở được tài liệu: dựng lại bằng python-docx
        # (đã là phụ thuộc sẵn có). Bản này mất ảnh và định dạng chữ, nhưng giữ
        # đúng thứ tự đoạn văn và bảng — đủ để đọc quy trình.
        return _docx_html_basic(path)
    with path.open("rb") as f:
        body = mammoth.convert_to_html(f).value
    return _header_html(path) + body


def _header_lines(path: Path) -> list[str]:
    """Chữ trong header của file, bỏ các trường tự động như số trang.

    mammoth bỏ qua header, mà phiếu thao tác của nhà máy lại để số phiếu ở
    header; quy trình thì để mã hiệu, lần ban hành. Thiếu phần này thì bản xem
    trên trình duyệt mất đúng thông tin nhận dạng tài liệu.
    """
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(str(path))
    section = doc.sections[0]
    headers = [section.header]
    if section.different_first_page_header_footer:
        headers.insert(0, section.first_page_header)

    lines: list[str] = []
    for header in headers:
        for p in header._element.iter(qn("w:p")):
            text, depth = [], 0
            for el in p.iter():
                # Trường phức (PAGE, NUMPAGES): chỉ còn giá trị lưu tạm lúc
                # soạn, hiển thị ra sẽ sai. Bỏ phần nằm giữa begin và end.
                if el.tag == qn("w:fldChar"):
                    kind = el.get(qn("w:fldCharType"))
                    depth += 1 if kind == "begin" else -1 if kind == "end" else 0
                elif el.tag == qn("w:t") and depth == 0:
                    parent = el.getparent()
                    inside_simple = parent is not None and parent.getparent() is not None \
                        and parent.getparent().tag == qn("w:fldSimple")
                    if not inside_simple:
                        text.append(el.text or "")
            line = "".join(text).strip()
            if line:
                lines.append(line)

    # Ghép nhãn đứng riêng ("Số phiếu:") với giá trị ở dòng sau; nhãn không có
    # giá trị (giá trị là trường số trang đã bỏ) thì bỏ luôn.
    merged: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.endswith(":"):
            nxt = lines[i + 1] if i + 1 < len(lines) else ""
            if nxt and not nxt.endswith(":"):
                merged.append(f"{line} {nxt}")
                i += 2
                continue
            i += 1
            continue
        merged.append(line)
        i += 1
    # Số trang vô nghĩa trên trình duyệt (cả tài liệu là một trang cuộn), và
    # phần còn sót của trường số trang như "Trang số: /5" chỉ gây rối.
    merged = [line for line in merged if not re.match(r"(?i)^trang(\s+số)?\b", line)]
    return list(dict.fromkeys(merged))


def _header_html(path: Path) -> str:
    try:
        lines = _header_lines(path)
    except Exception:
        return ""
    if not lines:
        return ""
    items = "".join(f"<div>{html.escape(line)}</div>" for line in lines)
    return f'<div class="doc-header">{items}</div>'


def _docx_html_basic(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    out: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            level = (para.style.name or "").lower()
            if level.startswith("heading"):
                out.append(f"<h3>{html.escape(text)}</h3>")
            else:
                out.append(f"<p>{html.escape(text)}</p>")
        elif tag == "tbl":
            out.append(_table_html(
                [[cell.text.strip() for cell in row.cells] for row in Table(child, doc).rows]
            ))
    return "\n".join(out)


def _xlsx_html(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    for sheet in wb.worksheets:
        rows = [
            ["" if v is None else str(v).strip() for v in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(r)]
        if rows:
            out.append(f"<h3>{html.escape(sheet.title)}</h3>")
            out.append(_table_html(rows))
    wb.close()
    return "\n".join(out)


def _csv_html(path: Path) -> str:
    import csv
    import io

    raw = path.read_text(encoding="utf-8", errors="replace")
    rows = [r for r in csv.reader(io.StringIO(raw)) if any(c.strip() for c in r)]
    return _table_html(rows)


def _table_html(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    head = "".join(f"<th>{html.escape(c)}</th>" for c in rows[0])
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>"
        for row in rows[1:]
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


# Khung trang: đặt bề rộng cột chữ vừa tầm đọc, bảng cuộn ngang được vì ma trận
# cắt của rơle rộng hơn màn hình.
_PAGE = """<!doctype html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{ color-scheme: light; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; background: #f1f5f9; color: #0f172a;
         font: 15px/1.65 "Segoe UI", system-ui, -apple-system, sans-serif; }}
  header {{ position: sticky; top: 0; z-index: 5; display: flex; flex-wrap: wrap;
            gap: 12px; align-items: center; justify-content: space-between;
            padding: 12px 20px; background: #0f766e; color: #fff;
            box-shadow: 0 1px 6px rgba(15, 23, 42, .25); }}
  header h1 {{ margin: 0; font-size: 16px; font-weight: 600; }}
  header .meta {{ font-size: 12px; opacity: .85; }}
  header a {{ color: #fff; font-size: 13px; text-decoration: none;
              border: 1px solid rgba(255, 255, 255, .55); border-radius: 6px;
              padding: 5px 12px; }}
  header a:hover {{ background: rgba(255, 255, 255, .15); }}
  main {{ max-width: 960px; margin: 24px auto; padding: 36px 44px; background: #fff;
          border-radius: 10px; box-shadow: 0 1px 3px rgba(15, 23, 42, .12); }}
  main img {{ max-width: 100%; height: auto; }}
  .doc-header {{ text-align: right; font-size: 13px; color: #334155; border-bottom: 1px solid #cbd5e1;
                 padding-bottom: 8px; margin-bottom: 18px; }}
  h1, h2, h3, h4 {{ line-height: 1.35; color: #0f172a; }}
  table {{ border-collapse: collapse; width: 100%; margin: 14px 0; font-size: 14px; }}
  th, td {{ border: 1px solid #cbd5e1; padding: 6px 9px; vertical-align: top;
            text-align: left; }}
  th {{ background: #f1f5f9; font-weight: 600; }}
  tr:nth-child(even) td {{ background: #f8fafc; }}
  /* Bảng rộng hơn khung thì cho cuộn ngang thay vì tràn ra ngoài trang. */
  main {{ overflow-x: auto; }}
  @media print {{
    body {{ background: #fff; }}
    header {{ display: none; }}
    main {{ box-shadow: none; margin: 0; padding: 0; max-width: none; }}
  }}
</style>
</head>
<body>
<header>
  <div>
    <h1>{title}</h1>
    <div class="meta">{meta}</div>
  </div>
  <div><a href="{download}">Tải tệp gốc</a></div>
</header>
<main>{body}</main>
</body>
</html>"""


def page(title: str, meta: str, body: str, download_url: str) -> str:
    return _PAGE.format(
        title=html.escape(title),
        meta=html.escape(meta),
        body=body,
        download=html.escape(download_url),
    )
