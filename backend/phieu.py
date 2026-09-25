"""Điền phiếu thao tác từ file Word mẫu do vận hành viên tải lên.

Vận hành viên giữ nguyên mẫu phiếu của nhà máy, chỉ gõ thêm các ô cần điền
dạng {{Tên ô}} vào đúng chỗ: {{Số phiếu}}, {{Người viết phiếu}}, {{Ngày}}...
Phần mềm tự đọc ra danh sách ô để dựng giao diện nhập liệu, rồi điền giá trị
vào đúng chỗ đó — định dạng, bảng biểu, logo, header của mẫu giữ nguyên.

Khó khăn chính: Word hay chẻ một đoạn chữ thành nhiều "run" (do sửa chính tả,
đổi phông giữa chừng, lưu nhiều lần). Phiếu thật của nhà máy có số phiếu
"053/2026" bị chẻ thành "053/", "202", "6". Nên không thể tìm-thay trên từng
run, mà phải ghép chữ của cả đoạn lại, tìm ô, rồi trả chữ về đúng các run.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .rag.textutils import normalize

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]{1,60}?)\s*\}\}")

# ------------------------------------------------------------------ loại ô

# Ô nhiều dòng: nội dung dài, nhập bằng ô soạn thảo thay vì một dòng.
_MULTILINE = ("noi dung", "ghi chu", "muc dich", "dieu kien", "luu y", "trinh tu",
              "su kien", "yeu cau", "bien phap")


def field_kind(name: str, all_names: list[str]) -> str:
    """Kiểu ô nhập liệu suy ra từ tên ô người dùng đặt trong mẫu.

    number: số phiếu, phần mềm tự cấp. date: ngày đầy đủ. day/month/year: mẫu
    viết "ngày {{Ngày}} tháng {{Tháng}} năm {{Năm}}", ba ô lấy chung một ngày.
    time: giờ. multiline / text: nhập tay.
    """
    key = normalize(name).strip()
    names = {normalize(n).strip() for n in all_names}
    if key in ("so phieu", "so") or key.startswith("so phieu"):
        return "number"
    split_date = "thang" in names and "nam" in names
    if split_date and key in ("ngay", "thang", "nam"):
        return {"ngay": "day", "thang": "month", "nam": "year"}[key]
    if key.startswith("ngay") or key.endswith(" ngay"):
        return "date"
    if key.startswith("gio") or " gio" in f" {key}":
        return "time"
    if any(word in key for word in _MULTILINE):
        return "multiline"
    return "text"


# ------------------------------------------------------- duyệt các phần file

def _parts(doc):
    """Mọi phần có chữ của file: thân, header, footer (kể cả trang đầu/chẵn)."""
    yield doc.element.body
    seen = set()
    for section in doc.sections:
        for hf in (section.header, section.first_page_header, section.even_page_header,
                   section.footer, section.first_page_footer, section.even_page_footer):
            # Header "liên kết với phần trước" trỏ về cùng một phần tử: bỏ qua
            # để không điền hai lần.
            if hf.is_linked_to_previous:
                continue
            el = hf._element
            if id(el) not in seen:
                seen.add(id(el))
                yield el


def _paragraphs(root):
    """Mọi đoạn văn trong một phần, gồm cả trong bảng lồng và khung văn bản."""
    return root.iter(qn("w:p"))


def _texts(p):
    """Các phần tử chữ thuộc riêng đoạn p (không lấy chữ của khung văn bản lồng)."""
    out = []
    for t in p.iter(qn("w:t")):
        owner = t.getparent()
        while owner is not None and owner.tag != qn("w:p"):
            owner = owner.getparent()
        if owner is p:
            out.append(t)
    return out


def _all_text(doc) -> str:
    return "\n".join(
        "".join(t.text or "" for t in _texts(p))
        for root in _parts(doc) for p in _paragraphs(root)
    )


# ------------------------------------------------------------------ đọc mẫu

def find_fields(path: Path) -> list[str]:
    """Tên các ô {{...}} trong mẫu, theo thứ tự xuất hiện, không trùng."""
    doc = Document(str(path))
    names: list[str] = []
    seen: set[str] = set()
    for match in PLACEHOLDER_RE.finditer(_all_text(doc)):
        name = unicodedata.normalize("NFC", match.group(1).strip())
        key = normalize(name)
        if key not in seen:
            seen.add(key)
            names.append(name)
    return names


# ------------------------------------------------------------------ điền mẫu

def _set_text(t, text: str) -> None:
    t.text = text
    # Không đặt thì Word nuốt dấu cách ở đầu/cuối chữ.
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _split_lines(t) -> None:
    """Chữ có xuống dòng thì tách thành chữ + ngắt dòng của Word."""
    lines = (t.text or "").split("\n")
    if len(lines) == 1:
        return
    _set_text(t, lines[0])
    anchor = t
    for line in lines[1:]:
        br = t.makeelement(qn("w:br"), {})
        anchor.addnext(br)
        new_t = t.makeelement(qn("w:t"), {})
        _set_text(new_t, line)
        br.addnext(new_t)
        anchor = new_t


def _replace_in_paragraph(p, pattern: re.Pattern, lookup) -> None:
    """Thay mọi chỗ khớp pattern trong đoạn p, kể cả khi trải qua nhiều run.

    lookup(match) trả về chữ thay thế, hoặc None để giữ nguyên.
    """
    texts = _texts(p)
    if not texts:
        return
    joined = "".join(t.text or "" for t in texts)
    # Vị trí bắt đầu của từng phần tử chữ trong chuỗi đã ghép.
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t.text or "")

    def locate(offset: int) -> int:
        idx = 0
        for i, s in enumerate(starts):
            if s <= offset:
                idx = i
        return idx

    touched = set()
    # Điền từ cuối lên để vị trí các ô phía trước không bị xê dịch.
    for match in reversed(list(pattern.finditer(joined))):
        value = lookup(match)
        if value is None:
            continue
        a, b = match.start(), match.end()
        first, last = locate(a), locate(b - 1)
        for i in range(first, last + 1):
            t = texts[i]
            text = t.text or ""
            lo = max(a - starts[i], 0)
            hi = min(b - starts[i], len(text))
            if i == first:
                _set_text(t, text[:lo] + value + text[hi:])
            else:
                _set_text(t, text[hi:])
            touched.add(i)
    for i in touched:
        _split_lines(texts[i])


def fill(path: Path, values: dict[str, str]) -> bytes:
    """Điền giá trị vào mẫu, trả về nội dung file .docx."""
    by_key = {normalize(k): v for k, v in values.items()}

    def lookup(match):
        return by_key.get(normalize(match.group(1).strip()))

    doc = Document(str(path))
    for root in _parts(doc):
        for p in list(_paragraphs(root)):
            _replace_in_paragraph(p, PLACEHOLDER_RE, lookup)
    return _save(doc)


def replace_text(path: Path, mapping: list[tuple[str, str]]) -> bytes:
    """Thay chữ cố định bằng chữ khác trên toàn file, chịu được run bị chẻ.

    Dùng để dựng mẫu từ một phiếu đã lập: thay tên người, ngày, số phiếu cụ
    thể bằng các ô {{...}}.
    """
    doc = Document(str(path))
    for old, new in mapping:
        pattern = re.compile(re.escape(old))
        for root in _parts(doc):
            for p in list(_paragraphs(root)):
                _replace_in_paragraph(p, pattern, lambda m, new=new: new)
    return _save(doc)


def _save(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ số phiếu

def format_number(pattern: str, number: int, year: int) -> str:
    """"###/YYYY/KH/HHC" → "053/2026/KH/HHC". Số dấu # là số chữ số tối thiểu."""
    pattern = pattern or "###/YYYY"
    out = re.sub(r"#+", lambda m: str(number).zfill(len(m.group(0))), pattern, count=1)
    if "#" not in pattern:
        out = f"{number}{pattern}" if pattern.startswith("/") else f"{pattern}{number}"
    return out.replace("YYYY", str(year)).replace("YY", str(year)[-2:])


def date_parts(value: str) -> dict[str, str]:
    """"2026-03-20" → các cách viết ngày cần cho mẫu."""
    try:
        d = date.fromisoformat(value)
    except (TypeError, ValueError):
        return {}
    return {
        "date": d.strftime("%d/%m/%Y"),
        "day": str(d.day),
        "month": str(d.month),
        "year": str(d.year),
    }
